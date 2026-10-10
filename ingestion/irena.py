"""IRENA: a manually downloaded CSV/Excel file copied unchanged into the raw layer.

Lands as irena_capacity.<ext> (extension kept). The original file name, size, modified time,
dataset version and download date go into the manifest.
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from ingestion.base import BaseIngestor, FetchResult, RawFile, SkipSource

log = logging.getLogger(__name__)

SUPPORTED = {".csv", ".xlsx", ".xls"}


class IRENAIngestor(BaseIngestor):
    source = "irena"
    file_glob = "irena_capacity.*"

    def fetch(self) -> FetchResult:
        cfg = self.settings.irena
        if not cfg.get("enabled", True):
            raise SkipSource("irena.enabled is false")

        configured = cfg.get("source_path") or "data/irena"
        path = Path(configured)
        if not path.is_absolute():
            path = self.settings.project_root / path
        source_file = self._pick_file(path)

        content = source_file.read_bytes()
        stat = source_file.stat()
        modified = datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc)
        ext = source_file.suffix.lower()

        details: dict[str, Any] = {
            "original_filename": source_file.name,
            "source_file_modified_utc": modified.isoformat(timespec="seconds"),
            "dataset_version": cfg.get("dataset_version"),
            "download_date": str(cfg.get("download_date") or modified.date().isoformat()),
            "download_date_is_file_mtime": not cfg.get("download_date"),
        }
        row_count, extra = _count_rows(content, ext)
        details.update(extra)
        if not cfg.get("dataset_version"):
            log.warning("[irena] irena.dataset_version is not set in config.yaml: the manifest will not record the file version")

        return FetchResult(
            files=[RawFile(name=f"irena_capacity{ext}", content=content, row_count=row_count)],
            request={"source": "manual download", "original_filename": source_file.name},
            details=details,
        )

    @staticmethod
    def _pick_file(path: Path) -> Path:
        if path.is_file():
            if path.suffix.lower() not in SUPPORTED:
                raise SkipSource(f"{path.name} is not a .csv/.xlsx/.xls file")
            return path
        if path.is_dir():
            files = [p for p in path.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED]
            if files:
                return max(files, key=lambda p: p.stat().st_mtime)
        raise SkipSource(
            f"no IRENA file found at '{path}'. Download it from IRENA (optional source) and put it in data/irena/"
        )


def _count_rows(content: bytes, ext: str) -> tuple[int | None, dict[str, Any]]:
    """Data rows (excluding the header). Informational only; failure never blocks ingestion."""
    try:
        if ext == ".csv":
            try:
                text = content.decode("utf-8-sig")
                encoding = "utf-8"
            except UnicodeDecodeError:
                text = content.decode("latin-1")
                encoding = "latin-1"
            rows = [r for r in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in r)]
            return max(len(rows) - 1, 0), {"encoding_guess": encoding}
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets = {ws.title: max(ws.max_row - 1, 0) for ws in wb.worksheets}
            return sum(sheets.values()), {"sheets": sheets}
    except Exception as exc:  # noqa: BLE001
        log.warning("[irena] could not count rows: %s", exc)
    return None, {}
